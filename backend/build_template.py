from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
title = doc.add_heading("COMPLIANCEFORGE", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_heading("Fiche de synthèse du traitement", level=2)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

table = doc.add_table(rows=6, cols=2)
table.style = "Light Grid Accent 1"
rows_data = [
    ("Organisation", "{{ organisation_nom }}"),
    ("Traitement", "{{ traitement_nom }}"),
    ("Service responsable", "{{ service_responsable }}"),
    ("Finalité", "{{ finalite }}"),
    ("Base légale", "{{ base_legale }}"),
    ("Statut", "{{ statut }}"),
]
for i, (label, value) in enumerate(rows_data):
    table.cell(i, 0).paragraphs[0].add_run(label).bold = True
    table.cell(i, 1).paragraphs[0].add_run(value)

doc.add_paragraph()
doc.add_heading("Catégories de données traitées", level=3)
cat_table = doc.add_table(rows=1, cols=2)
cat_table.style = "Light Grid Accent 1"
hdr = cat_table.rows[0].cells
hdr[0].paragraphs[0].add_run("Catégorie").bold = True
hdr[1].paragraphs[0].add_run("Sensibilité").bold = True
loop_row = cat_table.add_row().cells
loop_row[0].paragraphs[0].add_run("{%tr for cat in categories_donnees %}")
row2 = cat_table.add_row().cells
row2[0].paragraphs[0].add_run("{{ cat.categorie }}")
row2[1].paragraphs[0].add_run("{{ cat.sensibilite_label }}")
endloop_row = cat_table.add_row().cells
endloop_row[0].paragraphs[0].add_run("{%tr endfor %}")

doc.add_paragraph()
doc.add_heading("Scores", level=3)
score_table = doc.add_table(rows=2, cols=2)
score_table.style = "Light Grid Accent 1"
score_table.cell(0, 0).paragraphs[0].add_run("Score de conformité documentaire").bold = True
score_table.cell(0, 1).paragraphs[0].add_run("{{ score_conformite }} / 100")
score_table.cell(1, 0).paragraphs[0].add_run("Score de risque").bold = True
score_table.cell(1, 1).paragraphs[0].add_run("{{ score_risque }} / 100 ({{ niveau_risque }})")

doc.add_paragraph()
footer = doc.add_paragraph()
footer_run = footer.add_run(
    "Document généré automatiquement le {{ date_generation }} par ComplianceForge — "
    "ne constitue pas en l'état un document officiel APDP."
)
footer_run.italic = True
footer_run.font.size = Pt(9)
footer_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.save("templates/fiche_synthese_v1.docx")
print("Template créé : templates/fiche_synthese_v1.docx")
